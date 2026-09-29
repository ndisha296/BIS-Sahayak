"""Import BIS PDFs and catalogues into the local database and Chroma corpus."""
import argparse
import hashlib
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.database import (
    ProductStandardMap,
    SessionLocal,
    TestingFacility,
    init_db,
)
from app.rag import add_documents

_CHUNK_SIZE = 2400
_CHUNK_OVERLAP = 250
_BATCH_SIZE = 32


def _chunks(text: str):
    text = re.sub(r"\s+", " ", text).strip()
    start = 0
    while start < len(text):
        end = min(start + _CHUNK_SIZE, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + _CHUNK_SIZE // 2, end)
            if boundary > start:
                end = boundary
        yield text[start:end]
        if end == len(text):
            break
        start = max(end - _CHUNK_OVERLAP, start + 1)


def _stable_id(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized(value) -> str:
    return re.sub(r"[^a-z0-9]+", " ", _text(value).casefold()).strip()


def _header_columns(rows, required: dict[str, tuple[str, ...]]):
    for row_index, row in enumerate(rows[:30]):
        normalized = [_normalized(value) for value in row]
        columns = {}
        for field, names in required.items():
            columns[field] = next(
                (
                    index for index, cell in enumerate(normalized)
                    if any(name in cell for name in names)
                ),
                None,
            )
        if all(index is not None for index in columns.values()):
            return row_index, columns
    return None, None


def _workbook_rows(path: Path, required: dict[str, tuple[str, ...]]):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            header_row, columns = _header_columns(rows, required)
            if columns is None:
                continue
            for row in rows[header_row + 1:]:
                yield {
                    field: _text(row[index]) if index < len(row) else ""
                    for field, index in columns.items()
                }
    finally:
        workbook.close()


def _index_pdf(path: Path, root: Path, pending: list[tuple[str, str, dict]]) -> int:
    indexed_pages = 0
    relative = path.relative_to(root).as_posix()
    try:
        reader = PdfReader(str(path))
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                continue
            indexed_pages += 1
            for chunk_number, chunk in enumerate(_chunks(text), start=1):
                clause_match = re.search(
                    r"\b(?:clause|sub\s*clause)\s+([\d.]+[a-z]?)\b",
                    chunk,
                    re.IGNORECASE,
                )
                identity = f"pdf:{relative}:{page_number}:{chunk_number}"
                metadata = {
                    "source": relative,
                    "page": page_number,
                    "document_type": "standard_pdf",
                }
                if clause_match:
                    metadata["clause"] = clause_match.group(1)
                pending.append((_stable_id(identity), chunk, metadata))
    except Exception as error:
        print(f"Skipped unreadable PDF {relative}: {error}", file=sys.stderr)
    return indexed_pages


def _import_compulsory_products(path: Path, db: Session, pending) -> int:
    required = {
        "standard": ("is no", "indian standard no"),
        "product": ("product",),
    }
    imported = 0
    for row in _workbook_rows(path, required):
        standard = row["standard"]
        product = row["product"]
        if not standard or not product or not re.search(r"\bIS\b", standard, re.IGNORECASE):
            continue
        notification = row.get("notification", "")
        notes = "Compulsory certification catalogue. Confirm the current order and applicable scheme."
        if notification:
            notes = f"{notes} Notification: {notification}"
        exists = (
            db.query(ProductStandardMap)
            .filter(
                ProductStandardMap.product_keyword == product,
                ProductStandardMap.is_standard == standard,
            )
            .first()
        )
        if not exists:
            db.add(ProductStandardMap(
                product_keyword=product,
                is_standard=standard,
                scheme="Compulsory certification",
                notes=notes,
            ))
        identity = f"catalog:product:{standard}:{product}"
        pending.append((
            _stable_id(identity),
            f"Product: {product}. Indian Standard: {standard}. {notes}",
            {"source": path.name, "document_type": "product_catalogue"},
        ))
        imported += 1
    return imported


def _import_testing_facilities(path: Path, db: Session, pending) -> int:
    required = {
        "lab": ("lab name",),
        "standard": ("indian standard no",),
        "product": ("product",),
    }
    optional = {
        "osl": ("osl code",),
        "grade": ("grade type size designation",),
        "charges": ("testing charges",),
        "validity": ("validity date",),
        "remarks": ("remark",),
    }
    all_columns = {**required, **optional}
    imported = 0
    for row in _workbook_rows(path, all_columns):
        if not row["lab"] or not row["standard"]:
            continue
        exists = (
            db.query(TestingFacility)
            .filter(
                TestingFacility.lab_name == row["lab"],
                TestingFacility.indian_standard_no == row["standard"],
                TestingFacility.product == (row["product"] or None),
            )
            .first()
        )
        validity = row.get("validity", "")
        valid_until = None
        if isinstance(validity, (date, datetime)):
            valid_until = validity.date() if isinstance(validity, datetime) else validity
        if not exists:
            db.add(TestingFacility(
                lab_name=row["lab"],
                osl_code=row.get("osl") or None,
                indian_standard_no=row["standard"],
                product=row.get("product") or None,
                grade_type_size=row.get("grade") or None,
                testing_charges=row.get("charges") or None,
                validity_date=valid_until,
                remarks=row.get("remarks") or None,
            ))
        identity = f"catalog:lab:{row['lab']}:{row['standard']}:{row.get('product', '')}"
        pending.append((
            _stable_id(identity),
            "Testing laboratory: " + row["lab"]
            + f". Indian Standard: {row['standard']}. Product: {row.get('product', '')}."
            + f" Grade/type/size: {row.get('grade', '')}. Charges: {row.get('charges', '')}."
            + f" Validity: {row.get('validity', '')}. Remarks: {row.get('remarks', '')}",
            {"source": path.name, "document_type": "testing_facility_catalogue"},
        ))
        imported += 1
    return imported


def ingest(data_dir: Path) -> dict:
    if not data_dir.is_dir():
        raise ValueError(f"Data directory does not exist: {data_dir}")
    init_db()
    pending = []
    pdf_count = 0
    for pdf_path in sorted(data_dir.rglob("*.pdf")):
        pdf_count += _index_pdf(pdf_path, data_dir, pending)

    products_path = next(
        (path for path in data_dir.rglob("*.xlsx")
         if "products under compulsory certification" in path.stem.casefold()),
        None,
    )
    facilities_path = next(
        (path for path in data_dir.rglob("*.xlsx")
         if "testing facilities" in path.stem.casefold()),
        None,
    )

    db = SessionLocal()
    try:
        product_count = _import_compulsory_products(products_path, db, pending) if products_path else 0
        facility_count = _import_testing_facilities(facilities_path, db, pending) if facilities_path else 0
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

    for start in range(0, len(pending), _BATCH_SIZE):
        add_documents(pending[start:start + _BATCH_SIZE])
    return {
        "pdf_pages_indexed": pdf_count,
        "product_rows_imported": product_count,
        "testing_facility_rows_imported": facility_count,
        "rag_chunks_upserted": len(pending),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=Path("./data"))
    args = parser.parse_args()
    result = ingest(args.data_dir.expanduser().resolve())
    for label, count in result.items():
        print(f"{label}: {count}")


if __name__ == "__main__":
    main()